"""
Excel exports for Page Manager.

All three endpoints scope through visible_pages(request.user) — a non-supervisor
must never export a page (or a post inside a page) they don't own. The
"หมายเหตุ เฉพาะหัวหน้า" column is included only when the requesting user is
a supervisor.
"""
from urllib.parse import quote

import openpyxl
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from openpyxl.utils import get_column_letter

from apps.core.access import system_required

from .models import PagePost
from .scoping import is_supervisor, visible_pages
from .views import _apply_page_filters


def _xlsx_response(wb, filename):
    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    wb.save(resp)
    return resp


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 40)


def _thai_year_date(d):
    if not d:
        return ""
    return f"{d.day:02d}/{d.month:02d}/{d.year + 543}"


def _posts_workbook(posts, include_page_column, include_supervisor_note):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "โพสต์"

    header = ["ลำดับ"]
    if include_page_column:
        header.append("ชื่อเพจ")
    header += ["รูป", "ผู้ลงโพส", "รหัสสื่อ", "ประเภทสื่อ", "ID POST", "วันที่ลงสื่อ", "หมายเหตุ (ทั่วไปเฉพาะพนักงาน)"]
    if include_supervisor_note:
        header.append("ข้อมูลหมายเหตุ เฉพาะหัวหน้า")
    ws.append(header)

    for i, post in enumerate(posts, start=1):
        sku_label = f"[{post.product_code}] {post.product_name}".strip() if post.product_code else ""
        row = [i]
        if include_page_column:
            row.append(post.page.page_name)
        row += [
            post.image.url if post.image else "",
            post.poster.name if post.poster else "",
            sku_label,
            post.media_type.name if post.media_type else "",
            post.post_id,
            _thai_year_date(post.post_date),
            post.note,
        ]
        if include_supervisor_note:
            row.append(post.supervisor_note)
        ws.append(row)

    _autosize(ws)
    return wb


@system_required("pagemanager")
def export_page_posts(request, page_pk):
    page = get_object_or_404(visible_pages(request.user), pk=page_pk)
    sup = is_supervisor(request.user)
    posts = list(page.posts.select_related("poster", "media_type"))
    wb = _posts_workbook(posts, include_page_column=False, include_supervisor_note=sup)
    return _xlsx_response(wb, f"โพสต์ {page.page_name}.xlsx")


@system_required("pagemanager")
def export_all_posts(request):
    sup = is_supervisor(request.user)
    pages = visible_pages(request.user)
    posts = list(
        PagePost.objects.filter(page__in=pages).select_related("page", "poster", "media_type")
    )
    wb = _posts_workbook(posts, include_page_column=True, include_supervisor_note=sup)
    return _xlsx_response(wb, "โพสต์ทุกเพจ.xlsx")


@system_required("pagemanager")
def export_pages(request):
    qs = _apply_page_filters(visible_pages(request.user).prefetch_related("skus", "owners"), request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "เพจ"
    ws.append(["No.", "SKU", "ชื่อเพจ", "สถานะ", "ผู้ดูแล", "ID PAGE", "Link Page", "Facebook Chat"])

    for i, page in enumerate(qs, start=1):
        skus = ", ".join(s.product_code for s in page.skus.all())
        owners = ", ".join(o.name for o in page.owners.all())
        ws.append([i, skus, page.page_name, page.status, owners, page.page_id, page.page_url, page.chat_url])

    _autosize(ws)
    return _xlsx_response(wb, "รายการเพจ.xlsx")
